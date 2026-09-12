"""Prepare template-shaped result values from independently reconciled ledgers."""
from pathlib import Path
from datetime import date, timedelta
import argparse
import csv
import hashlib
import json
import sys
import numpy as np
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'src'))
from microgrid.validation import validate_absolute_run
from microgrid.settlement import validate_ledger
from microgrid.schemas import DispatchEvent, PenaltyEvent

ORIGIN=date(2025,1,1)
READ=lambda p:json.loads(p.read_text(encoding='utf-8'))


def main():
    ap=argparse.ArgumentParser()
    ap.add_argument('--out',required=True)
    args=ap.parse_args()
    out=Path(args.out).resolve()
    status=READ(out/'status.json')
    assert status['state'] in ('simulation_complete','complete')
    run_dir=Path(status['run_dir']);config=READ(run_dir/'config.json')
    problem=config['problem'];name=problem.replace('problem','result')+'.xlsx'
    template=ROOT/'data/附件5'/name
    raw=READ(run_dir/'summary.json');summary=raw['summary']
    z=dict(np.load(run_dir/'trajectory.npz'))
    first=date.fromisoformat(config['run_from']);last=date.fromisoformat(config['run_to'])
    assert first>ORIGIN, 'This result template reports after the January warm-up'
    start=(first-ORIGIN).days*1440+10;end=((last-ORIGIN).days+1)*1440+10
    np.testing.assert_array_equal(z['abs_minute'],np.arange(10,end,10))
    assert z['soc_boundary_kwh'][0]==6000 and not summary['infeasible_intervals']
    validate_absolute_run(abs_minutes=z['abs_minute'],grid_kwh=z['grid_kwh'],
        emergency_kwh=z['emergency_kwh'],charge_kwh=z['charge_kwh'],discharge_kwh=z['discharge_kwh'],
        curtail_kwh=z['curtail_kwh'],surplus_kwh=z['surplus_kwh'],soc_boundary_kwh=z['soc_boundary_kwh'],
        demand_kwh=z['demand_kwh'],pv_kwh=z['pv_kwh'],soc_start_kwh=6000,require_daily_cycle=False,allow_spill=True)
    ledger=READ(run_dir/'settlement_events.json');initial=READ(run_dir/'initial_plans.json')
    events=[DispatchEvent(**e) for e in ledger['execution']]
    penalties=[PenaltyEvent(**e) for e in ledger['penalties']]
    validate_ledger(abs_minutes=z['abs_minute'],grid_kwh=z['grid_kwh'],emergency_kwh=z['emergency_kwh'],
        price_actual=z['price_actual'],initial_plans=initial,events=events,penalties=penalties,
        can_adjust=problem in ('problem3','problem4-3'),expected_start_abs=10)
    np.testing.assert_allclose(z['plan_initial_kwh'],[initial[str(int(m))] for m in z['abs_minute']],atol=1e-7,rtol=0)
    np.testing.assert_allclose(z['plan_exec_kwh'],[e.o_exec_kwh for e in events],atol=1e-7,rtol=0)
    np.testing.assert_allclose(z['add_exec_kwh'],[e.a_exec_kwh for e in events],atol=1e-7,rtol=0)
    fees=z['price_actual']*(z['plan_exec_kwh']+1.5*z['add_exec_kwh']+5*z['emergency_kwh'])
    penalty_fees=np.zeros_like(fees)
    for p in penalties:
        penalty_fees[(p.at_abs-10)//10]+=.5*p.price_at_yuan_per_kwh*sum(p.reduced_kwh)
    fees+=penalty_fees
    bills={r['date']:r for r in raw['result_row_bills']}
    days=[first+timedelta(days=d) for d in range((last-first).days+1)]
    assert len(days)==summary['output_days']
    plan,adjust,charge,urgent=[],[],[],[]
    total_qty=total_fee=0.
    rnd=lambda x:round(float(x),6)
    def stamp(j):
        day,m=divmod((j+1)*10,1440)
        return f'{m//60}:{m%60:02d}'+('+1' if day else '')
    for day in days:
        base=(day-ORIGIN).days*1440;i=base//10
        # Array index zero is 00:10. Index base/10 is the date row's 00:10.
        sl=slice(i,i+144);bill=bills[day.isoformat()]
        quantity=float((z['grid_kwh'][sl]+z['emergency_kwh'][sl]).sum());fee=float(fees[sl].sum())
        assert abs(quantity-bill['total_kwh'])<1e-5 and abs(fee-bill['total_cost_yuan'])<1e-5
        plan.append([day.isoformat(),*[rnd(x) for x in z['plan_initial_kwh'][sl]],rnd(quantity),rnd(fee)])
        adjust.append([day.isoformat(),*[rnd(x) for x in z['grid_kwh'][sl]],rnd(quantity),rnd(fee)])
        total_qty+=quantity;total_fee+=fee
        # Four-hour blocks and SOC use the natural clock, independently of the
        # ten-minute shift of purchase rows. Template E/F rows 1 and 2 hold SOC.
        natural=base//10-1
        for b in range(6):
            a=natural+b*24;t=a+24
            time_label=0 if b==0 else '24:00' if b==1 else None
            soc=rnd(z['soc_boundary_kwh'][natural]) if b==0 else rnd(z['soc_boundary_kwh'][natural+144]) if b==1 else None
            charge.append([day.isoformat() if b==0 else None,f'{b*4}:00-{(b+1)*4}:00',
                rnd(z['charge_kwh'][a:t].sum()),rnd(z['discharge_kwh'][a:t].sum()),time_label,soc])
        hits=np.flatnonzero(z['emergency_kwh'][sl]>1e-9).tolist()
        groups=[]
        for j in hits:
            if groups and j==groups[-1][1] and j!=143:groups[-1][1]=j+1
            else:groups.append([j,j+1])
        for k,(a,b) in enumerate(groups):
            urgent.append([day.isoformat() if k==0 else None,f'{stamp(a)}-{stamp(b)}',rnd(z['emergency_kwh'][i+a:i+b].sum())])
    assert abs(total_qty-summary['total_purchased_kwh'])<1e-4
    assert abs(total_fee-summary['total_cost_yuan'])<1e-4
    body={'计划购电量':plan,'调整购电量':adjust,'充放电量':charge,'紧急购电量':urgent}
    source=load_workbook(template,read_only=False,data_only=False)
    sheets=[]
    for s in source:
        sheets.append({'name':s.title,'headers':[c.value for c in s[1]],'values':body[s.title],
            'template_rows':s.max_row,'columns':s.max_column,'body_height':s.row_dimensions[2].height or 14,
            'body_font':{'name':s['A2'].font.name,'size':s['A2'].font.sz,'bold':bool(s['A2'].font.b),
                         'italic':bool(s['A2'].font.i),'color':'#000000'}})
    source.close()
    hashes=READ(out/'source_and_template_hashes.json')
    assert all(hashlib.sha256(Path(p).read_bytes()).hexdigest()==h for p,h in hashes.items())
    payload={'problem':problem,'file_name':name,'template':str(template),'from':first.isoformat(),'to':last.isoformat(),
        'days':len(days),'execution_intervals':len(z['abs_minute']),'sheets':sheets,
        'report_quantity_kwh':total_qty,'report_cost_yuan':total_fee,'whole_run_cost_yuan':float(fees.sum()),
        'initial_soc_kwh':6000,'end_soc_kwh':float(z['soc_boundary_kwh'][-1]),
        'template_sha256':hashlib.sha256(template.read_bytes()).hexdigest()}
    (out/'result_payload.json').write_text(json.dumps(payload,ensure_ascii=False),encoding='utf-8')
    if problem.startswith('problem4'):
        # Separate test detail keeps the prescribed result workbook untouched.
        from datetime import datetime
        with (out/'逐时购电与费用.csv').open('w',encoding='utf-8-sig',newline='') as stream:
            writer=csv.writer(stream)
            writer.writerow(['开始时间','日初计划_kWh','保留原计划_kWh','调整增购_kWh','调整后普通购电_kWh',
                '紧急购电_kWh','实际外网总购电_kWh','实际电价_元每kWh','原计划执行费_元','调整增购费_元',
                '紧急购电费_元','本时刻调减违约费_元','本段总费用_元','电费标签','段初SOC_kWh','段末SOC_kWh','弃购电_kWh'])
            for i in np.flatnonzero((z['abs_minute']>=start)&(z['abs_minute']<end)):
                price=z['price_actual'][i]
                costs=[price*z['plan_exec_kwh'][i],1.5*price*z['add_exec_kwh'][i],5*price*z['emergency_kwh'][i]]
                writer.writerow([(datetime(2025,1,1)+timedelta(minutes=int(z['abs_minute'][i]))).isoformat(),
                    z['plan_initial_kwh'][i],z['plan_exec_kwh'][i],z['add_exec_kwh'][i],z['grid_kwh'][i],
                    z['emergency_kwh'][i],z['grid_kwh'][i]+z['emergency_kwh'][i],price,*costs,
                    penalty_fees[i],fees[i],f'{fees[i]:.2f}元',z['soc_boundary_kwh'][i],z['soc_boundary_kwh'][i+1],z['surplus_kwh'][i]])
    (out/'核验').mkdir(exist_ok=True)
    checks={k:v for k,v in payload.items() if k!='sheets'}
    checks.update(physics_valid=True,ledger_valid=True,all_daily_totals_reconciled=True,
        source_and_templates_unchanged=True)
    (out/'核验/result_source_checks.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(checks,ensure_ascii=True,indent=2))


if __name__=='__main__':
    main()
