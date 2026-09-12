"""Read-only verification of template structure and every filled result cell."""
from pathlib import Path
from datetime import datetime, date, time, timedelta
import argparse
import csv
import hashlib
import json
import numpy as np
from openpyxl import load_workbook


def main():
    ap=argparse.ArgumentParser();ap.add_argument('--out',required=True)
    out=Path(ap.parse_args().out).resolve()
    read=lambda p:json.loads(p.read_text(encoding='utf-8'))
    p=read(out/'result_payload.json');status=read(out/'status.json')
    file=out/p['file_name'];template=Path(p['template'])
    result=load_workbook(file,read_only=False,data_only=True)
    source=load_workbook(template,read_only=False,data_only=False)
    assert result.sheetnames==source.sheetnames==[s['name'] for s in p['sheets']]
    count=0
    def close(a,b,label,tol=1e-6):
        assert isinstance(a,(int,float)) and np.isfinite(a) and abs(a-b)<tol,(label,a,b)
    for item in p['sheets']:
        name=item['name'];s=result[name];original=source[name];cols=item['columns']
        assert [s.cell(1,c).value for c in range(1,cols+1)]==item['headers']
        assert set(map(str,s.merged_cells.ranges))==set(map(str,original.merged_cells.ranges))
        assert s.freeze_panes==original.freeze_panes
        assert len(s.tables)==len(original.tables)==0
        for c in range(1,cols+1):
            actual=s.cell(1,c);ref=original.cell(1,c)
            assert actual.font.name==ref.font.name and actual.font.sz==ref.font.sz,(name,c,'header font')
            assert actual.font.b==ref.font.b and actual.alignment.horizontal==ref.alignment.horizontal
        for i,row in enumerate(item['values'],2):
            for c,value in enumerate(row,1):
                cell=s.cell(i,c);a=cell.value;count+=1
                if c==1 and value is not None:
                    assert isinstance(a,datetime) and a==datetime.fromisoformat(value),(name,i,c,a,value)
                elif name=='充放电量' and c==5 and value==0:
                    assert a==time(0,0) or a==0,(name,i,c,a)
                elif value is None:
                    assert a is None,(name,i,c,a)
                elif isinstance(value,(int,float)):
                    close(a,value,(name,i,c))
                else:
                    assert a==value,(name,i,c,a,value)
                assert cell.data_type!='e'
        # Uncomputed template dates/placeholder rows stay empty, never zero-filled.
        for row in s.iter_rows(min_row=len(item['values'])+2,max_row=s.max_row,max_col=cols):
            assert all(c.value is None for c in row),(name,'unexpected trailing data')
    run_dir=Path(status['run_dir']);z=dict(np.load(run_dir/'trajectory.npz'))
    index={int(m):i for i,m in enumerate(z['abs_minute'])}
    first=date.fromisoformat(p['from']);origin=date(2025,1,1)
    # Independent minute arithmetic checks all columns, the next-day tail and
    # the SOC rows at 00:00 and 24:00 in the original six-row template block.
    for d in range(p['days']):
        day=first+timedelta(days=d);midnight=(day-origin).days*1440
        for j in range(144):
            pos=index[midnight+(j+1)*10]
            close(result['计划购电量'].cell(d+2,j+2).value,z['plan_initial_kwh'][pos],('plan clock',d,j))
            if '调整购电量' in result:
                close(result['调整购电量'].cell(d+2,j+2).value,z['grid_kwh'][pos],('adjust clock',d,j))
        s=result['充放电量'];r=d*6+2
        close(s.cell(r,6).value,z['soc_boundary_kwh'][index[midnight]],('midnight SOC',d))
        close(s.cell(r+1,6).value,z['soc_boundary_kwh'][index[midnight+1440]],('24h SOC',d))
        for b in range(6):
            positions=[index[midnight+(b*24+j)*10] for j in range(24)]
            close(s.cell(r+b,3).value,z['charge_kwh'][positions].sum(),('natural charge block',d,b))
            close(s.cell(r+b,4).value,z['discharge_kwh'][positions].sum(),('natural discharge block',d,b))
    for name in ('计划购电量','调整购电量'):
        if name not in result:continue
        close(sum(result[name].cell(r,146).value for r in range(2,p['days']+2)),p['report_quantity_kwh'],(name,'quantity total'),1e-3)
        close(sum(result[name].cell(r,147).value for r in range(2,p['days']+2)),p['report_cost_yuan'],(name,'fee total'),1e-3)
    csv_path=out/'逐时购电与费用.csv'
    if p['problem'].startswith('problem4'):
        with csv_path.open(encoding='utf-8-sig',newline='') as stream:
            details=list(csv.DictReader(stream))
        assert len(details)==p['days']*144
        amount=0.
        columns={'日初计划_kWh':'plan_initial_kwh','保留原计划_kWh':'plan_exec_kwh',
            '调整增购_kWh':'add_exec_kwh','调整后普通购电_kWh':'grid_kwh','紧急购电_kWh':'emergency_kwh',
            '实际电价_元每kWh':'price_actual','弃购电_kWh':'surplus_kwh'}
        start=datetime.combine(first,time(0,10))
        for i,row in enumerate(details):
            stamp=start+timedelta(minutes=i*10)
            assert datetime.fromisoformat(row['开始时间'])==stamp
            minute=int((stamp-datetime(2025,1,1)).total_seconds()/60);pos=index[minute]
            for key,array in columns.items():close(float(row[key]),z[array][pos],('detail',i,key))
            close(float(row['段初SOC_kWh']),z['soc_boundary_kwh'][pos],('detail start SOC',i))
            close(float(row['段末SOC_kWh']),z['soc_boundary_kwh'][pos+1],('detail end SOC',i))
            fee=float(row['本段总费用_元']);amount+=fee
            component=sum(float(row[k]) for k in ('原计划执行费_元','调整增购费_元','紧急购电费_元','本时刻调减违约费_元'))
            close(fee,component,('detail fee',i))
            assert float(row['本时刻调减违约费_元'])>=0
            assert row['电费标签'].endswith('元')
            close(float(row['电费标签'][:-1]),fee,('detail label',i),.00501)
        close(amount,p['report_cost_yuan'],'detail total fee',1e-3)
    hashes=read(out/'source_and_template_hashes.json')
    assert all(hashlib.sha256(Path(path).read_bytes()).hexdigest()==h for path,h in hashes.items())
    assert hashlib.sha256(template.read_bytes()).hexdigest()==p['template_sha256']
    checks={'file':str(file),'days':p['days'],'filled_cells_checked':count,
        'template_sheet_order_headers_and_header_styles_preserved':True,'all_values_and_dates_match':True,
        'independent_time_mapping_valid':True,'natural_day_soc_positions_preserved':True,
        'daily_and_period_totals_reconciled':True,'unused_rows_empty':True,'source_and_templates_unchanged':True,
        'problem4_detail_csv_verified':p['problem'].startswith('problem4'),
        'sha256':hashlib.sha256(file.read_bytes()).hexdigest()}
    (out/'核验/saved_result_checks.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2),encoding='utf-8')
    result.close();source.close();print(json.dumps(checks,ensure_ascii=True,indent=2))


if __name__=='__main__':
    main()
