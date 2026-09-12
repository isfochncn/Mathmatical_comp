"""Read-only checks of the saved problem-three spreadsheet against its source rows."""
from pathlib import Path
from datetime import datetime
import argparse
import hashlib
import json
from openpyxl import load_workbook


def main():
    ap=argparse.ArgumentParser()
    ap.add_argument('--out',required=True)
    out=Path(ap.parse_args().out)
    p=json.loads((out/'month_payload.json').read_text(encoding='utf-8'))
    file=out/'问题三_2025年2月_逐时购电与电费.xlsx'
    w=load_workbook(file,read_only=True,data_only=True)
    f=load_workbook(file,read_only=False,data_only=False)
    assert w.sheetnames==['月度汇总','逐时购电与费用','每日汇总','调减违约明细','供电与预测核验']
    def close(a,b,label,tol=1e-5):
        assert isinstance(a,(int,float)) and abs(a-b)<=tol,(label,a,b)
    def same_time(a,b,label):
        expected=datetime.fromisoformat(b)
        assert isinstance(a,datetime) and abs((a-expected).total_seconds())<.001,(label,a,b)
    for name,key,date_cols,text_cols in [('逐时购电与费用','rows',3,{20}),
                                       ('供电与预测核验','check_rows',2,set()),
                                       ('调减违约明细','penalty_rows',2,set())]:
        source=p[key]
        if not source:
            continue
        count=len(source);cols=len(source[0])
        assert f[name].max_row==count+4 and f[name].max_column==cols
        assert f[name].freeze_panes=='C5' and len(f[name].tables)==1
        for i,row in enumerate(w[name].iter_rows(min_row=5,max_row=count+4,max_col=cols,values_only=True)):
            for c in range(cols):
                if c<date_cols:
                    same_time(row[c],source[i][c],(name,i,c))
                elif c in text_cols:
                    assert isinstance(row[c],str) and row[c].endswith('元')
                    close(float(row[c][:-1]),source[i][19],('fee label',i),.00501)
                else:
                    close(row[c],source[i][c],(name,i,c))
    keys=['plan_kwh','retained_kwh','add_kwh','ordinary_kwh','net_adjustment_kwh','emergency_kwh',
          'total_kwh','plan_cost','add_cost','emergency_cost','reduced_kwh','penalty_cost','total_cost',
          'spill_kwh','soc_start','soc_end','emergency_intervals','revised_intervals','penalty_nodes','spill_intervals']
    for i,key in enumerate(keys):
        close(w['月度汇总'].cell(i+5,2).value,p['current'][key],key,1e-4)
    # Daily reconciliation checks each complete row, including the shifted dates
    # and shared SOC boundaries, rather than only the monthly grand total.
    daily=w['每日汇总']
    for d,row in enumerate(daily.iter_rows(min_row=5,max_row=p['n']//144+4,max_col=15,values_only=True)):
        source=p['rows'][d*144:(d+1)*144]
        same_time(row[0],source[0][2],('daily date',d))
        for c,raw in enumerate([4,5,6,7,9,10,12,14,16,18,19,21],1):
            close(row[c],sum(r[raw] for r in source),('daily',d,c),1e-4)
        close(row[13],p['check_rows'][d*144][11],('daily initial SOC',d))
        close(row[14],p['check_rows'][(d+1)*144-1][12],('daily final SOC',d))
    for row in (5,40,p['n']+4):
        assert f['逐时购电与费用'].cell(row,20).value==f'=M{row}+O{row}+Q{row}+S{row}'
        assert f['逐时购电与费用'].cell(row,8).value==f'=F{row}+G{row}'
    # Verify that every node's fee references its actual event records.
    for i,row in enumerate(p['rows']):
        span=p['penalty_ranges'].get(row[0])
        if span:
            a,b=span
            assert f['逐时购电与费用'].cell(i+5,19).value==f"=SUM('调减违约明细'!F{a}:F{b})"
    errors=[]
    for sheet in w:
        for row in sheet:
            errors.extend(f'{sheet.title}!{c.coordinate}:{c.value}' for c in row if c.data_type=='e')
    assert not errors,errors[:20]
    hashes=json.loads((out/'source_and_template_hashes.json').read_text(encoding='utf-8'))
    assert all(hashlib.sha256(Path(path).read_bytes()).hexdigest()==digest for path,digest in hashes.items())
    result={'month':p['month'],'n_intervals':p['n'],'from':p['from'],'to':p['to'],
        'saved_values_and_timestamps_match':True,'daily_totals_match':True,'four_fee_components_match':True,
        'penalty_attribution_matches':True,'source_and_templates_unchanged':True,'formula_errors':errors,
        'max_bus_residual_kwh':max(abs(r[20]) for r in p['check_rows']),
        'max_soc_residual_kwh':max(abs(r[21]) for r in p['check_rows']),
        'xlsx_sha256':hashlib.sha256(file.read_bytes()).hexdigest(),'workbook':str(file.resolve())}
    (out/'核验/saved_workbook_checks.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    w.close();f.close()
    print(json.dumps(result,ensure_ascii=True,indent=2))


if __name__=='__main__':
    main()
