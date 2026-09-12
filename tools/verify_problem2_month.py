"""Read-only verification of the saved monthly inspection workbook and caches."""
from pathlib import Path
from datetime import datetime
import argparse
import hashlib
import json
from openpyxl import load_workbook


def main():
    ap=argparse.ArgumentParser()
    ap.add_argument('--out', required=True)
    args=ap.parse_args()
    out=Path(args.out)
    payload=json.loads((out/'month_payload.json').read_text(encoding='utf-8'))
    year,month=map(int,payload['month'].split('-'))
    path=out/f'问题二_{year}年{month}月_逐时购电与电费.xlsx'
    values=load_workbook(path,read_only=True,data_only=True)
    formulas=load_workbook(path,read_only=False,data_only=False)
    assert values.sheetnames == ['月度汇总','逐时购电与费用','每日汇总','供电与预测核验']
    def close(a,b,label,tolerance=1e-5):
        assert isinstance(a,(float,int)) and abs(a-b)<=tolerance,(label,a,b)
    n=payload['n']; last=n+4
    detail=values['逐时购电与费用']
    # The exporter omits optional XML dimension metadata; read-only dimensions
    # can therefore be None. Check the parsed cells rather than that metadata.
    detail_cells=formulas['逐时购电与费用']
    assert detail_cells.max_row == last and detail_cells.max_column == 17
    fees=[]
    for i,row in enumerate(detail.iter_rows(min_row=5,max_row=last,values_only=True)):
        source=payload['rows'][i]
        for col in range(3):
            expected=datetime.fromisoformat(source[col])
            assert isinstance(row[col],datetime) and abs((row[col]-expected).total_seconds())<.001
        for col in list(range(3,14))+[15,16]:
            close(row[col],source[col],f'purchase row {i} column {col}')
        assert isinstance(row[14],str) and row[14].endswith('元')
        close(float(row[14][:-1]),row[13],f'fee label {i}',.00501)
        fees.append(row[13])
    physics=values['供电与预测核验']
    max_bus=max_stock=0.
    for i,row in enumerate(physics.iter_rows(min_row=5,max_row=last,values_only=True)):
        for col in range(2):
            expected=datetime.fromisoformat(payload['check_rows'][i][col])
            assert isinstance(row[col],datetime) and abs((row[col]-expected).total_seconds())<.001
        for col in range(2,22):
            close(row[col],payload['check_rows'][i][col],f'physics row {i} column {col}')
        max_bus=max(max_bus,abs(row[20]));max_stock=max(max_stock,abs(row[21]))
    keys=['plan_kwh','ordinary_kwh','emergency_kwh','total_kwh','ordinary_cost','emergency_cost','total_cost',
          'spill_kwh','soc_start','soc_end','emergency_intervals','revised_intervals']
    summary=values['月度汇总']
    for i,key in enumerate(keys):
        close(summary.cell(5+i,2).value,payload['current'][key],key,1e-4)
        close(summary.cell(5+i,3).value,payload['legacy'][key],f'legacy {key}',1e-4)
    close(sum(fees),payload['current']['total_cost'],'sum of interval fees',1e-4)
    daily=values['每日汇总']
    close(sum(row[0] for row in daily.iter_rows(min_row=5,max_row=n//144+4,min_col=8,max_col=8,values_only=True)),
          payload['current']['total_cost'],'daily/monthly fee reconciliation',1e-4)
    for name in ('逐时购电与费用','供电与预测核验'):
        sheet=formulas[name]
        assert sheet.freeze_panes == 'C5',(name,sheet.freeze_panes)
        assert len(sheet.tables) == 1
    for row in (5,148,last):
        assert formulas['逐时购电与费用'].cell(row,14).value == f'=K{row}+M{row}'
        assert formulas['逐时购电与费用'].cell(row,7).value == f'=F{row}-E{row}'
    errors=[]
    for sheet in values:
        for row in sheet:
            errors.extend(f'{sheet.title}!{cell.coordinate}:{cell.value}' for cell in row if cell.data_type=='e')
    assert not errors,errors[:20]
    hashes=json.loads((out/'source_and_template_hashes.json').read_text(encoding='utf-8'))
    assert all(hashlib.sha256(Path(p).read_bytes()).hexdigest()==h for p,h in hashes.items())
    result={'month':payload['month'],'n_intervals':n,'from':payload['from'],'to':payload['to'],
        'purchase_and_fee_values_match':True,'physics_values_match':True,'cached_formulas_match':True,
        'formula_errors':errors,'max_bus_residual_kwh':max_bus,'max_soc_residual_kwh':max_stock,
        'source_and_templates_unchanged':True,'xlsx_sha256':hashlib.sha256(path.read_bytes()).hexdigest(),
        'workbook':str(path.resolve())}
    (out/'核验/saved_workbook_checks.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    values.close();formulas.close()
    print(json.dumps(result,ensure_ascii=True,indent=2))


if __name__=='__main__':
    main()
