"""Regression checks against main-model contracts and independent small examples."""
from copy import deepcopy
from dataclasses import replace
from datetime import date, timedelta
from types import SimpleNamespace

import numpy as np
import pytest

from microgrid.data_io import Attachment1, Attachment2, Attachment3, Attachment4, DataBundle, ForecastBlock
from microgrid.forecast import Forecaster
from microgrid.timeline import build_timeline
from microgrid.planning import FeeMode, WindowForecast, solve_window, clear_window_cache, window_cache_info
from microgrid.simulation import execute_interval, InfeasibleInterval, RunOptions, run_absolute
from microgrid.absolute_run import POLICIES, RunConfig, run, save_run
from microgrid.validation import validate_absolute_run, ValidationError


@pytest.fixture
def synthetic():
    days = [date(2025, 1, 1) + timedelta(days=i) for i in range(365)]
    return DataBundle(
        Attachment1(np.ones(144), np.full(144, 600.0), np.zeros(144)),
        Attachment2(days, np.full((365, 144), 600.0), np.zeros((365, 144))),
        Attachment3(tuple(ForecastBlock(d, h, np.zeros(24)) for d in days for h in (0, 6, 12, 18))),
        Attachment4(days, np.full((365, 144), 3.0)),
    )


def forecast(demand, price=1):
    d = np.asarray(demand, dtype=float)
    return WindowForecast(np.arange(len(d))*10, d, np.zeros(len(d)),
                          np.full(len(d), price, dtype=float), np.full(len(d), "test"))


def test_exact_adjustment_premium_and_next_day_fee():
    r = solve_window(forecast=forecast([10, 10]), soc_start_kwh=1200,
                     fee_mode=FeeMode.ADJUSTABLE, o_kwh=np.zeros(2), a_kwh=np.zeros(2),
                     adjustable_mask=np.array([True, False]), price_now_yuan_per_kwh=1)
    assert r.report.feasible
    assert r.grid_kwh == pytest.approx([10, 10])
    assert r.objective_yuan == pytest.approx(15+10)


def test_zero_objective_is_zero_and_cache_refreshes_inputs():
    clear_window_cache()
    r = solve_window(forecast=forecast([10], 0), soc_start_kwh=1200, fee_mode=FeeMode.FIRST_PLAN)
    assert r.objective_yuan == 0
    r = solve_window(forecast=forecast([20], 2), soc_start_kwh=1200, fee_mode=FeeMode.FIRST_PLAN)
    assert r.objective_yuan == pytest.approx(40)
    assert window_cache_info()["hit"] == 1


def test_no_emergency_in_problem_one():
    r = solve_window(forecast=forecast([1000]), soc_start_kwh=1200,
                     fee_mode=FeeMode.FROZEN, o_kwh=np.zeros(1), a_kwh=np.zeros(1),
                     allow_emergency=False)
    assert not r.report.feasible
    assert np.isnan(r.grid_kwh).all()
    with pytest.raises(RuntimeError):
        r.require_ok()


def test_feedback_keeps_simultaneous_targets_and_real_losses():
    a = execute_interval(abs_minute=0, grid_committed_kwh=19, demand_kwh=0, pv_kwh=0,
                         soc_now_kwh=10800, charge_target_kwh=90, discharge_target_kwh=81, allow_spill=False)
    assert [a.charge_kwh, a.discharge_kwh, a.soc_end_kwh] == pytest.approx([90, 81, 10800])
    assert a.spill_kwh == 0


def test_feedback_curtails_only_pv_and_refuses_impossible_overpurchase():
    a = execute_interval(abs_minute=0, grid_committed_kwh=0, demand_kwh=0, pv_kwh=100,
                         soc_now_kwh=10800, charge_target_kwh=0, discharge_target_kwh=0)
    assert a.curtail_kwh == pytest.approx(100)
    with pytest.raises(InfeasibleInterval):
        execute_interval(abs_minute=0, grid_committed_kwh=1000, demand_kwh=0, pv_kwh=0,
                         soc_now_kwh=10800, charge_target_kwh=0, discharge_target_kwh=0, allow_spill=False)


def test_soc_compatibility_conversion():
    a = execute_interval(abs_minute=0, grid_committed_kwh=0, demand_kwh=100, pv_kwh=0,
                         soc_now_kwh=6000, soc_target_kwh=5900, allow_spill=False)
    assert a.discharge_kwh == pytest.approx(90)
    assert a.soc_end_kwh == pytest.approx(5900)


def test_grid_direct_supply_is_not_limited_by_battery_charging():
    r = solve_window(forecast=forecast([2000]), soc_start_kwh=6000,
                     soc_end_fixed_kwh=6000, fee_mode=FeeMode.FIRST_PLAN,
                     allow_emergency=False).require_ok()
    assert r.grid_kwh == pytest.approx([2000])
    assert r.charge_kwh == pytest.approx([0])
    assert r.discharge_kwh == pytest.approx([0])
    a = execute_interval(abs_minute=10, grid_committed_kwh=2000, demand_kwh=2000,
                         pv_kwh=0, soc_now_kwh=6000, charge_target_kwh=0, discharge_target_kwh=0)
    assert a.soc_end_kwh == pytest.approx(6000)
    assert a.charge_kwh == pytest.approx(0)


@pytest.mark.parametrize("soc,grid", [(6000, 1384.8989), (10800, 800)])
def test_purchase_surplus_cannot_exceed_power_or_capacity(soc, grid):
    # Demand is 550.9598: the first case exceeds charging power by 0.605767.
    # At full SOC the second exceeds even the permitted simultaneous losses.
    r = solve_window(forecast=forecast([550.9598]), soc_start_kwh=soc,
                     fee_mode=FeeMode.FROZEN, o_kwh=np.array([grid]), a_kwh=np.zeros(1))
    assert not r.report.feasible
    with pytest.raises(InfeasibleInterval):
        execute_interval(abs_minute=10, grid_committed_kwh=grid, demand_kwh=550.9598,
                         pv_kwh=0, soc_now_kwh=soc, charge_target_kwh=0, discharge_target_kwh=0, allow_spill=False)


def test_cheap_grid_purchase_can_charge_for_expensive_period():
    f = forecast([0, 750])
    f.price_yuan_per_kwh = np.array([1., 3.])
    r = solve_window(forecast=f, soc_start_kwh=6000, soc_end_fixed_kwh=6000,
                     fee_mode=FeeMode.FIRST_PLAN, allow_emergency=False).require_ok()
    assert r.charge_kwh == pytest.approx([750, 0])
    assert r.discharge_kwh == pytest.approx([0, 675])
    assert r.grid_kwh == pytest.approx([750/.9, 75])
    assert r.objective_yuan < 750*3  # cheaper than direct grid supply only


def test_historical_pv_kw_is_converted_once(synthetic):
    synthetic.attachment2.pv_actual_kw[:] = 600
    f = Forecaster(synthetic, build_timeline(synthetic))
    assert f.pv_forecast_kwh(1440, np.array([1440]), use_published=False).values[0] == pytest.approx(100)


def test_future_actual_mutation_cannot_change_forecast_or_bridge(synthetic):
    f1 = Forecaster(synthetic, build_timeline(synthetic))
    before = f1.window_forecast(0, 0, 2880, use_published_pv=False, current_price=None)
    changed = deepcopy(synthetic)
    changed.attachment2.load_kw[:] = 99999
    changed.attachment2.pv_actual_kw[:] = 99999
    changed.attachment4.price_yuan_per_kwh[:] = 999
    f2 = Forecaster(changed, build_timeline(changed))
    after = f2.window_forecast(0, 0, 2880, use_published_pv=False, current_price=None)
    for name in ("demand_kwh", "pv_kwh", "price_yuan_per_kwh"):
        assert np.array_equal(getattr(before, name), getattr(after, name))
    assert np.isnan(f1.timeline.load_kw.values[0]) and np.isnan(f2.timeline.load_kw.values[0])


def test_problem_two_never_calls_attachment_three(synthetic):
    class Forbidden:
        def __getattr__(self, name):
            raise AssertionError("attachment3 accessed")
    b = replace(synthetic, attachment3=Forbidden())
    f = Forecaster(b, build_timeline(b))
    f.window_forecast(0, 0, 2880, use_published_pv=False, current_price=None)


def test_midnight_includes_yesterday_and_recent_fallback(synthetic):
    tl = build_timeline(synthetic)
    tl.load_kw.values[:144] = 120
    f = Forecaster(synthetic, tl)
    samples, n = f._same_clock_samples(1440, tl.load_kw.values)
    assert n == 1
    assert np.nanmean(samples[:, 100]) == 120
    # A future clock without same-clock history falls back to mature observations.
    p = f.point_forecast(30, tl.load_kw.values, np.array([100]), "load",
                         cold_start_kw=np.full(144, 999))
    assert p.values[0] == 120


def test_mature_history_uses_latest_28_slots_without_future_leak(synthetic):
    synthetic.attachment2.load_kw[:] = np.arange(365)[:, None]*6 + 600
    now = 30*1440 + 720
    tl = build_timeline(synthetic)
    f = Forecaster(synthetic, tl)
    samples, _ = f._same_clock_samples(now, tl.load_kw.values)
    # 11:50 has completed today; 12:00 has not. Both use exactly 28 mature days.
    assert np.sum(np.isfinite(samples[:, 71])) == 28
    assert np.sum(np.isfinite(samples[:, 72])) == 28
    assert np.nanmean(samples[:, 71]) == pytest.approx(600 + np.mean(np.arange(3, 31))*6)
    assert np.nanmean(samples[:, 72]) == pytest.approx(600 + np.mean(np.arange(2, 30))*6)
    before = f.window_forecast(now, now, 32*1440, use_published_pv=False, current_price=None)
    for series in (tl.load_kw, tl.pv_kw, tl.price_yuan_per_kwh):
        series.values[now//10:] = 99999
    after = f.window_forecast(now, now, 32*1440, use_published_pv=False, current_price=None)
    for name in ("demand_kwh", "pv_kwh", "price_yuan_per_kwh"):
        assert np.array_equal(getattr(before, name), getattr(after, name))


def test_source_tail_retained_and_unobserved_extension_missing(synthetic):
    synthetic.attachment2.load_kw[-1, -1] = 1234
    tl = build_timeline(synthetic)
    assert tl.load_kw.value_at(365*1440) == 1234
    with pytest.raises(Exception, match="unavailable"):
        tl.load_kw.value_at(365*1440+10)
    f = Forecaster(synthetic, tl)
    r = f.window_forecast(365*1440, 365*1440, 367*1440,
                         use_published_pv=True, current_price=None)
    assert len(r.abs_minutes) == 288
    assert np.isfinite(r.demand_kwh).all()


@pytest.mark.parametrize("problem", ["problem2", "problem3", "problem4-2", "problem4-3"])
def test_rolling_permissions_prices_and_full_horizon(synthetic, problem):
    tl = build_timeline(synthetic)
    f = Forecaster(synthetic, tl)
    result = run_absolute(timeline=tl, forecaster=f, policy=POLICIES[problem],
                          options=RunOptions(), abs_from=10, abs_to=370, soc_start_kwh=6000)
    assert result.window_solves == 36
    assert result.forecast_audits[0]["formed_at"] == 10
    assert all(a["window_end"] == 2880 for a in result.forecast_audits)
    expected_price = 1 if problem in ("problem2", "problem3") else 3
    assert result.run.abs_minutes[0] == 10
    assert all(e.price_actual_yuan_per_kwh == expected_price for e in result.run.events)
    assert len(result.run.initial_plans) == 143
    assert 0 not in result.run.initial_plans
    assert max(result.run.initial_plans) == 1430
    assert not result.released_abs
    assert all(s.surplus_kwh == 0 for s in result.run.steps)


def test_end_to_end_problem_one_save_and_export(synthetic, tmp_path):
    synthetic.attachment1.load_kw[:] = np.arange(144)*6+600
    r = run(RunConfig(problem="problem1", out_dir=tmp_path), synthetic)
    assert r.export_arrays["demand_kwh"][0] == synthetic.attachment1.demand_kwh[-1]
    assert r.summary["day_balance_residual_kwh"] == pytest.approx(0, abs=1e-6)
    path = save_run(r)
    from microgrid.cli import build_parser, cmd_export
    assert cmd_export(build_parser().parse_args(["export", "--problem", "problem1", "--out", str(tmp_path)])) == 0
    from openpyxl import load_workbook
    wb = load_workbook(path / "result/result1.xlsx", read_only=True, data_only=True)
    assert wb["计划购电量"].cell(2, 2).value == pytest.approx(r.export_arrays["result_grid_kwh"][0], abs=1e-6)
    wb.close()


def test_main_config_rejects_hidden_policy_changes():
    assert RunConfig().plan_refresh_intervals == 1
    assert RunConfig().allow_spill
    for kw in (dict(plan_refresh_intervals=6), dict(absorption_safety_kwh=600),
               dict(warmup_days=1), dict(max_infeasible_intervals=1)):
        with pytest.raises(ValueError):
            RunConfig(**kw)


def test_validation_rejects_disposal_and_nonfinite_data():
    args = dict(abs_minutes=np.array([0]), grid_kwh=np.array([10.]),
                emergency_kwh=np.zeros(1), charge_kwh=np.zeros(1), discharge_kwh=np.zeros(1),
                curtail_kwh=np.zeros(1), surplus_kwh=np.array([10.]), soc_boundary_kwh=np.array([6000., 6000.]),
                demand_kwh=np.zeros(1), pv_kwh=np.zeros(1), soc_start_kwh=6000)
    with pytest.raises(ValidationError):
        validate_absolute_run(**args)

    args["grid_kwh"][:] = np.nan
    with pytest.raises(ValidationError):
        validate_absolute_run(**args)


def test_revision_archive_penalty_and_export_round_trip(synthetic, monkeypatch, tmp_path):
    """Controlled plans isolate accounting/export from the optimizer's choice."""
    from microgrid import simulation, constants
    from microgrid.cli import build_parser, cmd_export
    from microgrid.export import build_daily_rows
    import json

    def controlled_solver(**kw):
        f = kw["forecast"]
        if kw["fee_mode"] == FeeMode.FROZEN:
            grid = np.where(kw["committed_mask"], kw["committed_grid_kwh"], 100.)
        else:
            grid = np.full(f.n, 100. if kw["fee_mode"] == FeeMode.FIRST_PLAN else 80.)
        result = SimpleNamespace(grid_kwh=grid, charge_kwh=np.zeros(f.n),
                                 discharge_kwh=np.maximum(100-grid, 0), objective_yuan=0.,
                                 reserve_shortfall_kwh=0., risk_penalty_yuan=0.)
        result.reserve_stock_shortfall_kwh = 0.
        result.reserve_power_shortfall_kwh = np.zeros(f.n)
        result.require_ok = lambda: result
        return result

    monkeypatch.setattr(simulation, "solve_window", controlled_solver)
    monkeypatch.setattr(constants, "OUTPUT_START", (2025, 1, 1))
    r = run(RunConfig(problem="problem4-3", pv_method="pooled", run_to=date(2025, 1, 1), out_dir=tmp_path), synthetic)
    assert len(r.run.steps) == 144  # includes the result row's next-midnight actual
    assert np.array_equal(r.run.abs_minutes, np.arange(10, 1450, 10))
    assert len(r.run.soc_boundary_kwh) == 145
    assert r.run.soc_boundary_kwh[0] == 6000
    assert 0 not in r.run.initial_plans
    assert r.summary["warmup_intervals"] == 0
    assert r.summary["output_soc_start_abs_minute"] == 10
    assert r.export_arrays["demand_kwh"][0] == synthetic.attachment2.load_kw[0, 0]/6
    penalty = r.run.penalties[0]
    assert penalty.at_abs == 360
    assert penalty.price_at_yuan_per_kwh == 3  # actual, not forecast or prior quote
    assert penalty.penalty_yuan == pytest.approx(.5*3*20*108)
    assert r.run.initial_plans[360] == 100
    assert r.run.events[35].o_exec_kwh == 80
    assert r.forecast_audits[-1]["window_end"] == 4320
    rows = build_daily_rows(r)
    assert rows[0].plan_initial_kwh[35] == 100
    assert rows[0].final_plan_kwh[35] == 80
    assert rows[0].natural_start_minute == 10
    assert rows[0].charge_clock_start_kwh is None
    from microgrid.export import paper_table2
    assert paper_table2(rows)[0]["时间段"] == "0:10-4:00"
    assert paper_table2(rows)[6]["时间段"] == "0:10储电量"
    assert r.run.soc_boundary_kwh[143] == pytest.approx(rows[0].soc_natural_end_kwh)
    path = save_run(r)
    saved = json.loads((path / "summary.json").read_text(encoding="utf-8"))
    assert saved["natural_day_bills"][0]["from_abs"] == 10
    assert saved["natural_day_bills"][0]["to_abs"] == 1440
    assert not saved["natural_day_bills"][0]["complete_natural_day"]
    args = build_parser().parse_args(["export", "--problem", "problem4-3", "--out", str(tmp_path)])
    assert cmd_export(args) == 0
    from openpyxl import load_workbook
    wb = load_workbook(path / "result/result4-3.xlsx", read_only=True, data_only=True)
    assert wb["计划购电量"].cell(2, 37).value == 100
    assert wb["调整购电量"].cell(2, 37).value == 80
    assert wb["充放电量"].cell(2, 2).value == "0:10-4:00"
    assert wb["充放电量"].cell(2, 5).value == "0:10"
    assert wb["充放电量"].cell(2, 6).value == 6000
    wb.close()
    # The saved ledger is revalidated, rather than trusting a validation stamp.
    ledger_path = path / "settlement_events.json"
    ledger = json.loads(ledger_path.read_text())
    ledger["penalties"][0]["price_at_yuan_per_kwh"] = 1
    ledger_path.write_text(json.dumps(ledger))
    with pytest.raises(Exception, match="actual price"):
        cmd_export(args)
    later = run(RunConfig(problem="problem4-3", pv_method="pooled", run_from=date(2025, 1, 2),
                          run_to=date(2025, 1, 2)), synthetic)
    assert later.run.abs_minutes[0] == 10
    assert [d for d, _ in later.row_bills] == [date(2025, 1, 2)]
    assert later.summary["soc_at_output_start_kwh"] == pytest.approx(3600, abs=1e-5)
    assert later.summary["warmup_intervals"] == 143


def test_paper_charge_blocks_use_natural_midnight():
    from microgrid.export import DailyExportRow, paper_table2
    row = DailyExportRow(date(2025, 2, 1), *[np.zeros(144) for _ in range(4)],
                         np.arange(1., 145.), np.arange(1., 145.), 6000, 6000, 0, 0, 999, 888)
    blocks = paper_table2([row])
    assert blocks[0]["充电量"] == 999+sum(range(1, 24))
    assert blocks[5]["充电量"] == sum(range(120, 144))


@pytest.mark.parametrize("soc,expected_spill,expected_charge", [(6000, 1000-750/.9, 750), (10800, 1000, 0)])
def test_unavoidable_paid_spill_power_vs_capacity(soc, expected_spill, expected_charge):
    r = solve_window(forecast=forecast([0], 2), soc_start_kwh=soc, fee_mode=FeeMode.FROZEN,
                     o_kwh=np.array([1000.]), a_kwh=np.zeros(1), allow_spill=True).require_ok()
    assert r.spill_kwh[0] == pytest.approx(expected_spill, abs=1e-5)
    assert r.objective_yuan == pytest.approx(2000, abs=1e-5)  # no refund on disposal
    a = execute_interval(abs_minute=10, grid_committed_kwh=1000, demand_kwh=0,
                         pv_kwh=50, soc_now_kwh=soc, charge_target_kwh=0, discharge_target_kwh=750)
    assert a.spill_kwh == pytest.approx(expected_spill, abs=1e-6)
    assert a.charge_kwh == pytest.approx(expected_charge, abs=1e-6)
    assert a.curtail_kwh == pytest.approx(50)
    assert a.discharge_kwh == 0  # no artificial burning to reduce the disposal metric


def test_new_purchase_plan_prefers_stored_energy_and_no_disposal():
    r = solve_window(forecast=forecast([100, 100]), soc_start_kwh=1200+100/.9,
                     fee_mode=FeeMode.FIRST_PLAN, allow_spill=True).require_ok()
    assert r.grid_kwh == pytest.approx([0, 100], abs=1e-4)
    assert r.spill_kwh == pytest.approx([0, 0])
    assert r.discharge_kwh[0] == pytest.approx(100, abs=1e-4)


def test_adjustable_plan_reduces_overpurchase_instead_of_discarding():
    r = solve_window(forecast=forecast([100]), soc_start_kwh=10800, fee_mode=FeeMode.ADJUSTABLE,
                     o_kwh=np.array([1000.]), a_kwh=np.zeros(1),
                     price_now_yuan_per_kwh=1, allow_spill=True).require_ok()
    assert r.spill_kwh == pytest.approx([0])
    assert r.grid_kwh[0] <= 100+1e-5


def test_spill_validation_keeps_paid_quantity_and_rejects_avoidable_waste():
    args = dict(abs_minutes=np.array([10]), grid_kwh=np.array([1000.]), emergency_kwh=np.zeros(1),
                charge_kwh=np.zeros(1), discharge_kwh=np.zeros(1), curtail_kwh=np.zeros(1),
                surplus_kwh=np.array([1000.]), soc_boundary_kwh=np.array([10800., 10800.]),
                demand_kwh=np.zeros(1), pv_kwh=np.zeros(1), soc_start_kwh=10800)
    assert validate_absolute_run(**args).total_surplus_kwh == 1000
    args["soc_boundary_kwh"][:] = 6000
    args["soc_start_kwh"] = 6000
    with pytest.raises(ValidationError, match="Avoidable"):
        validate_absolute_run(**args)


def test_disposal_cache_refreshes_headroom_and_respects_strict_mode():
    clear_window_cache()
    args = dict(soc_start_kwh=10800, fee_mode=FeeMode.FROZEN, o_kwh=np.array([1000.]), a_kwh=np.zeros(1))
    r = solve_window(forecast=forecast([0]), allow_spill=True, **args).require_ok()
    assert r.spill_kwh[0] == pytest.approx(1000)
    r = solve_window(forecast=forecast([1000]), allow_spill=True, **args).require_ok()
    assert r.spill_kwh[0] == pytest.approx(0, abs=2e-6)
    assert not solve_window(forecast=forecast([0]), allow_spill=False, **args).report.feasible
